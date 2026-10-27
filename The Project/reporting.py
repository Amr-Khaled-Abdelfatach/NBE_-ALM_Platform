"""Custom PDF and PPTX Executive Report Generator with NBE Header, Footer, Logo, Scenario Parameters, and AI Insights."""
import io
import zipfile
from datetime import datetime
import pandas as pd

from config import NBE_GREEN, NBE_DARK_GREEN, NBE_ORANGE, PARAM_SPECS, SCENARIO_DEFAULTS, OLLAMA_HOST, OLLAMA_MODEL
from scenarios import ask_ollama

REPORTING_AVAILABLE = True
REPORTING_MISSING_ERROR = ""

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak
    )
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError as exc:
    REPORTING_AVAILABLE = False
    REPORTING_MISSING_ERROR = str(exc)


def _get_ollama_executive_recommendations(scenario_name: str, params: dict, selected_kpis: list) -> str:
    """يستغل نموذج Ollama لتوليد ملخص تنفيذي وتوصيات إستراتيجية للجنة ALCO."""
    try:
        prompt_payload = {
            "task": "Generate an Executive Summary, Key Risks, and Strategic ALCO Recommendations for NBE based on this ALM scenario.",
            "scenario_name": scenario_name,
            "parameters": params,
            "key_performance_indicators": selected_kpis
        }
        response = ask_ollama(OLLAMA_HOST, OLLAMA_MODEL, prompt_payload)
        if response and not str(response).startswith("Phi-3 error"):
            return str(response).replace("**", "")
    except Exception:
        pass
    return "• Executive Summary: Scenario evaluated successfully under deterministic stress framework.\n• Recommendations: Maintain optimal liquidity buffers and monitor interest rate repricing gaps closely."


def _hex_to_rgb_tuple(hex_str: str):
    h = hex_str.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def plotly_fig_to_png_bytes(fig, width=900, height=420) -> bytes:
    fig.update_layout(paper_bgcolor="#FFFFFF", plot_bgcolor="#FFFFFF", font=dict(color="#173B2A"))
    return fig.to_image(format="png", width=width, height=height, scale=2)


def _add_pdf_header_footer(canvas, doc):
    """Draws professional Header (with Logo) and Footer on every page of the PDF."""
    canvas.saveState()
    
    try:
        canvas.drawImage("nbe_logo.png", landscape(A4)[0] - 120, landscape(A4)[1] - 32, width=80, height=25, preserveAspectRatio=True, mask='auto')
    except Exception:
        pass

    canvas.setFont("Helvetica-Bold", 8)
    canvas.setFillColor(colors.HexColor(NBE_DARK_GREEN))
    
    canvas.drawString(36, landscape(A4)[1] - 25, "NATIONAL BANK OF EGYPT (NBE) — ALM SIMULATION CENTER")
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawRightString(landscape(A4)[0] - 130, landscape(A4)[1] - 25, f"Generated: {datetime.now():%Y-%m-%d %H:%M}")
    
    canvas.setStrokeColor(colors.HexColor(NBE_GREEN))
    canvas.setLineWidth(1)
    canvas.line(36, landscape(A4)[1] - 35, landscape(A4)[0] - 36, landscape(A4)[1] - 35)
    
    canvas.line(36, 35, landscape(A4)[0] - 36, 35)
    canvas.drawString(36, 22, "CONFIDENTIAL — For Internal Management & ALCO Use Only")
    canvas.drawRightString(landscape(A4)[0] - 36, 22, f"Page {doc.page}")
    canvas.restoreState()


def generate_pdf_report(
    scenario_name: str, scenario_type: str, simulation_date: str, horizon_days: int,
    selected_kpis: list, selected_charts: dict, comparison_df: pd.DataFrame = None,
    risk_df: pd.DataFrame = None, ai_insight: str = None, params: dict = None,
    engine=None, run_id=None
) -> bytes:
    if not REPORTING_AVAILABLE:
        raise RuntimeError(f"Missing reporting libraries ({REPORTING_MISSING_ERROR}). Run: pip install reportlab python-pptx kaleido")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=0.5 * inch, rightMargin=0.5 * inch,
        topMargin=0.6 * inch, bottomMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("NBE_Title", parent=styles["Heading1"], fontSize=16, textColor=colors.HexColor(NBE_DARK_GREEN), spaceAfter=4)
    subtitle_style = ParagraphStyle("NBE_Subtitle", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#555555"), spaceAfter=10)
    section_style = ParagraphStyle("NBE_Section", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor(NBE_GREEN), spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle("NBE_Body", parent=styles["Normal"], fontSize=9, leading=13)

    chip_style = ParagraphStyle(
        "NBE_Chip", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor(NBE_DARK_GREEN),
        backColor=colors.HexColor("#EAF5EF"), borderPadding=4, borderWidth=0.5,
        borderColor=colors.HexColor(NBE_GREEN), spaceAfter=4, spaceRight=4
    )

    story = [
        Paragraph(f"Executive ALM Simulation Report: {scenario_name}", title_style),
        Paragraph(f"<b>Scenario Type:</b> {scenario_type} &nbsp;|&nbsp; <b>Simulation Date:</b> {simulation_date} &nbsp;|&nbsp; <b>Horizon:</b> {horizon_days} Days", subtitle_style),
    ]

    if params:
        story.append(Paragraph("Active Scenario Parameters & Shocks", section_style))
        chip_texts = []
        for k, v in params.items():
            if k in PARAM_SPECS:
                label, unit = PARAM_SPECS[k][0], PARAM_SPECS[k][1]
                val_str = f"{float(v):+g}{unit}" if isinstance(v, (int, float)) else f"{v}{unit}"
                chip_texts.append(f"<b>{label}:</b> {val_str}")
        
        if chip_texts:
            chip_table_data = []
            row = []
            for i, text in enumerate(chip_texts):
                row.append(Paragraph(text, chip_style))
                if len(row) == 4 or i == len(chip_texts) - 1:
                    while len(row) < 4:
                        row.append("")
                    chip_table_data.append(row)
                    row = []
            
            ct = Table(chip_table_data, colWidths=[2.0 * inch, 2.0 * inch, 2.0 * inch, 2.0 * inch])
            ct.setStyle(TableStyle([
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]))
            story.append(ct)
            story.append(Spacer(1, 6))

    # --- منطق البحث عن التفسير في القاعدة أو توليده وحفظه تلقائياً ---
    final_insight = ai_insight
    if not final_insight and engine is not None and run_id is not None:
        try:
            df_run = pd.read_sql(f"SELECT ai_explanation FROM simulation_runs WHERE run_id = {run_id}", engine)
            if not df_run.empty and pd.notna(df_run.iloc[0]["ai_explanation"]):
                final_insight = df_run.iloc[0]["ai_explanation"]
        except Exception:
            pass

    if not final_insight:
        final_insight = _get_ollama_executive_recommendations(scenario_name, params, selected_kpis)
        if engine is not None and run_id is not None and not str(final_insight).startswith("Phi-3 error"):
            try:
                from database import update_run_ai_explanation
                update_run_ai_explanation(engine, run_id, str(final_insight))
            except Exception:
                pass

    if final_insight:
        story.append(Paragraph("Scenario Explanation & Management Insight", section_style))
        clean_insight = str(final_insight).replace("**", "")
        story.append(Paragraph(clean_insight.replace("\n", "<br/>"), body_style))
        story.append(Spacer(1, 8))

    if selected_kpis:
        story.append(Paragraph("Selected Key Performance Indicators (KPIs)", section_style))
        kpi_rows = [["KPI Metric", "Baseline (Before)", "Scenario (After)", "Impact / Delta"]]
        for item in selected_kpis:
            kpi_rows.append([str(item["label"]), str(item["before"]), str(item["after"]), str(item["delta"] or "-")])
        t = Table(kpi_rows, colWidths=[2.8 * inch, 2.0 * inch, 2.0 * inch, 2.0 * inch])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NBE_GREEN)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D7E6DD")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EAF5EF")]),
        ]))
        story.append(t)
        story.append(Spacer(1, 8))

    if selected_charts:
        story.append(Paragraph("Selected Simulation Charts", section_style))
        for chart_title, fig in selected_charts.items():
            try:
                img_bytes = plotly_fig_to_png_bytes(fig, width=900, height=380)
                story.append(Paragraph(f"<b>{chart_title}</b>", body_style))
                story.append(Spacer(1, 2))
                story.append(RLImage(io.BytesIO(img_bytes), width=7.8 * inch, height=3.2 * inch))
                story.append(Spacer(1, 10))
            except Exception as exc:
                story.append(Paragraph(f"Chart error ({chart_title}): {exc}", body_style))

    if comparison_df is not None and not comparison_df.empty:
        story.append(PageBreak())
        story.append(Paragraph("Detailed Scenario Comparison Table", section_style))
        df_clean = comparison_df.drop(columns=[c for c in ["Source", "source"] if c in comparison_df.columns], errors="ignore")
        comp_data = [list(df_clean.columns)]
        for _, r in df_clean.iterrows():
            comp_data.append([f"{v:,.2f}" if isinstance(v, (int, float)) and pd.notna(v) else str(v) for v in r.values])
        ct_table = Table(comp_data, repeatRows=1)
        ct_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NBE_DARK_GREEN)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FBF9")]),
        ]))
        story.append(ct_table)

    doc.build(story, onFirstPage=_add_pdf_header_footer, onLaterPages=_add_pdf_header_footer)
    return buffer.getvalue()


def generate_pptx_report(
    scenario_name: str, scenario_type: str, simulation_date: str, horizon_days: int,
    selected_kpis: list, selected_charts: dict, comparison_df: pd.DataFrame = None,
    ai_insight: str = None, params: dict = None,
    engine=None, run_id=None
) -> bytes:
    if not REPORTING_AVAILABLE:
        raise RuntimeError(f"Missing reporting libraries ({REPORTING_MISSING_ERROR}). Run: pip install reportlab python-pptx kaleido")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    nbe_green = RGBColor(*_hex_to_rgb_tuple(NBE_GREEN))
    nbe_dark = RGBColor(*_hex_to_rgb_tuple(NBE_DARK_GREEN))
    nbe_orange = RGBColor(*_hex_to_rgb_tuple(NBE_ORANGE))

    def add_nbe_slide_header_footer(slide, title, sub=""):
        banner = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
        banner.fill.solid()
        banner.fill.fore_color.rgb = nbe_dark
        banner.line.fill.background()
        tf = banner.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        if sub:
            p2 = tf.add_paragraph()
            p2.text = sub
            p2.font.size = Pt(9)
            p2.font.color.rgb = nbe_orange

        try:
            slide.shapes.add_picture("nbe_logo_w.png", Inches(11.8), Inches(0.15), height=Inches(0.6))
        except Exception:
            pass

        footer = slide.shapes.add_shape(1, Inches(0), Inches(7.1), Inches(13.333), Inches(0.4))
        footer.fill.solid()
        footer.fill.fore_color.rgb = RGBColor(240, 245, 242)
        footer.line.color.rgb = RGBColor(200, 220, 210)
        ftf = footer.text_frame
        fp = ftf.paragraphs[0]
        fp.text = "National Bank of Egypt (NBE) — ALM Simulation Center | Confidential (ALCO Use Only)"
        fp.font.size = Pt(8)
        fp.font.color.rgb = RGBColor(100, 120, 110)

    # --- Slide 1: KPIs ---
    slide1 = prs.slides.add_slide(blank_layout)
    add_nbe_slide_header_footer(slide1, f"NBE ALM Simulation Report — {scenario_name}", f"Type: {scenario_type} | Date: {simulation_date} | Horizon: {horizon_days} Days")
    if selected_kpis:
        rows = len(selected_kpis) + 1
        ts = slide1.shapes.add_table(rows, 4, Inches(0.6), Inches(1.2), Inches(12.1), Inches(0.38 * rows))
        t = ts.table
        headers = ["KPI Metric", "Baseline (Before)", "Scenario (After)", "Impact / Delta"]
        for c, h in enumerate(headers):
            cell = t.cell(0, c)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = nbe_green
            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                p.font.size = Pt(11)
                p.font.color.rgb = RGBColor(255, 255, 255)
        for r, item in enumerate(selected_kpis, start=1):
            for c, val in enumerate([item["label"], item["before"], item["after"], item["delta"] or "-"]):
                cell = t.cell(r, c)
                cell.text = str(val)
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(10)

    # --- Slide 2: Scenario Parameters ---
    if params:
        slide_params = prs.slides.add_slide(blank_layout)
        add_nbe_slide_header_footer(slide_params, f"Applied Scenario Parameters — {scenario_name}", "Active macroeconomic & market stress drivers")
        
        left_pos, top_pos = Inches(0.8), Inches(1.5)
        chip_width, chip_height = Inches(3.6), Inches(0.6)
        h_gap, v_gap = Inches(0.3), Inches(0.2)
        col_limit = 3

        idx = 0
        for k, v in params.items():
            if k in PARAM_SPECS:
                label, unit = PARAM_SPECS[k][0], PARAM_SPECS[k][1]
                val_str = f"{float(v):+g}{unit}" if isinstance(v, (int, float)) else f"{v}{unit}"
                
                c_x = left_pos + (idx % col_limit) * (chip_width + h_gap)
                c_y = top_pos + (idx // col_limit) * (chip_height + v_gap)
                
                shape = slide_params.shapes.add_shape(5, c_x, c_y, chip_width, chip_height)
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(234, 245, 239)
                shape.line.color.rgb = RGBColor(0, 102, 68)
                
                tf = shape.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = f"{label}: {val_str}"
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = RGBColor(23, 59, 42)
                
                idx += 1

    # --- البحث عن التفسير للبوربوينت أو توليده وحفظه ---
    final_insight = ai_insight
    if not final_insight and engine is not None and run_id is not None:
        try:
            df_run = pd.read_sql(f"SELECT ai_explanation FROM simulation_runs WHERE run_id = {run_id}", engine)
            if not df_run.empty and pd.notna(df_run.iloc[0]["ai_explanation"]):
                final_insight = df_run.iloc[0]["ai_explanation"]
        except Exception:
            pass

    if not final_insight:
        final_insight = _get_ollama_executive_recommendations(scenario_name, params, selected_kpis)
        if engine is not None and run_id is not None and not str(final_insight).startswith("Phi-3 error"):
            try:
                from database import update_run_ai_explanation
                update_run_ai_explanation(engine, run_id, str(final_insight))
            except Exception:
                pass

    # --- Slide 3: AI / Management Insight ---
    if final_insight:
        slide_ai = prs.slides.add_slide(blank_layout)
        add_nbe_slide_header_footer(slide_ai, "AI Scenario Explanation & Management Insight", f"Scenario: {scenario_name}")
        tb = slide_ai.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(11.9), Inches(5.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = str(final_insight).replace("**", "")
        tf.paragraphs[0].font.size = Pt(13)

    # --- Slides 4+: Selected Charts ---
    for chart_title, fig in selected_charts.items():
        slide_c = prs.slides.add_slide(blank_layout)
        add_nbe_slide_header_footer(slide_c, chart_title, f"Scenario: {scenario_name} ({simulation_date})")
        try:
            png_bytes = plotly_fig_to_png_bytes(fig, width=1100, height=480)
            slide_c.shapes.add_picture(io.BytesIO(png_bytes), Inches(0.8), Inches(1.2), width=Inches(11.7))
        except Exception as exc:
            tb = slide_c.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(2))
            tb.text_frame.text = f"Chart error ({chart_title}): {exc}"

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def generate_zip_bundle(pdf_bytes: bytes, pptx_bytes: bytes, base_filename: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{base_filename}.pdf", pdf_bytes)
        zf.writestr(f"{base_filename}.pptx", pptx_bytes)
    return buf.getvalue()